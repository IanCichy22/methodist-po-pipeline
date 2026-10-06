"""
End-to-end MAXISMED purchase-order pipeline, wrapped as a single class:
pulls PO-request emails from Gmail, redeems their Office 365 Message
Encryption attachments over Playwright, merges the extracted SeqNo/PO
records into Fulfilled/Unfulfilled JSON files (new data wins on duplicate
keys), and backfills matching PO numbers -- plus duplicate/discrepancy
flags -- into the "Invoice Line Items" sheet of the commission tracker.

Gmail is driven entirely over IMAP with an app password -- Google blocks
Playwright/Chromium from logging into the Gmail web UI ("This browser or app
may not be secure"), and app passwords aren't valid on the Gmail login form
anyway (protocol access only). Playwright is used ONLY for the Microsoft OME
portal (outlook.office365.com/Encryption/...), which authenticates the
recipient via emailed one-time passcode and needs no Google session at all.
The "Read the message" link is extracted directly from the raw email HTML
rather than clicked through Gmail's UI.
"""

import csv
import email
import imaplib
import json
import re
import time
from pathlib import Path
from urllib.parse import urlparse

import openpyxl
import pandas as pd
from bs4 import BeautifulSoup
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
from playwright.sync_api import sync_playwright

READ_LINK_PHRASES = (
    "view the message",
    "read the message",
    "view message",
    "read message",
    "open the message",
    "open message",
    "view encrypted message",
    "read encrypted message",
    "view secure message",
    "read secure message",
)
# Words within a phrase are joined with \s+ so &nbsp; / stray whitespace between
# words in the source HTML doesn't break the match. \b boundaries matter: without
# them, "...Microsoft Purview Message Encryption" false-matches "view message"
# (Pur-VIEW + Message) -- that boilerplate footer link is in every OME
# notification, so this isn't a hypothetical.
READ_LINK_PATTERN = re.compile(
    "|".join(
        r"\b" + r"\s+".join(re.escape(w) for w in phrase.split()) + r"\b" for phrase in READ_LINK_PHRASES
    ),
    re.IGNORECASE,
)

# The genuine "read the message" link always redeems through an office365.com
# host; requiring that filters out any text-pattern false positive (like the
# Purview footer link, which points at go.microsoft.com) and any quoted/duplicate
# link from an older notification that happens to carry matching anchor text.
READ_LINK_HREF_HOST_SUFFIXES = ("office365.com",)

# Matches "one-time passcode", "One Time Passcode", "one-time code", etc.
# Deliberately does NOT match "passkey" -- that's a different (WebAuthn) sign-in
# option, not a wording variant of the email-OTP flow this class automates.
OTP_LABEL_RE = re.compile(r"one[\s-]?time.*?code", re.I)

# The "check your email" screen has several elements matching OTP_LABEL_RE
# (heading, instructions, field label, "get another one-time passcode" button)
# so a bare get_by_text(OTP_LABEL_RE) is a strict-mode violation there. This
# narrows to the one-off "We sent a ... passcode" heading to confirm that
# specific screen has gone away.
OTP_SENT_NOTICE_RE = re.compile(r"we sent.*one[\s-]?time.*code", re.I)


class PurchaseOrderPipeline:
    """
    Runs the whole MAXISMED PO workflow end to end: fetch_emails() pulls
    unread PO-request emails and parses their attachments; merge_and_split()
    combines new records into the Fulfilled/Unfulfilled JSON files on disk;
    backfill_tracker() writes matching PO numbers (and duplicate/discrepancy
    flags) into the commission tracker. run() chains all three.

    Each stage can also be called independently -- e.g. backfill_tracker()
    needs no email credentials, so it's the one to call repeatedly while
    iterating on tracker behavior without re-pulling email each time.

    Usage:
        pipeline = PurchaseOrderPipeline(
            email_user="you@gmail.com",
            gmail_app_password="....",
            sender_filter="someone@example.com",
            subject_filter="Purchase Order Number Request",
            po_column_name="Purchase Order #",
            key_columns=["Facility", "DOS", ...],
            fulfilled_json_path="output/Fulfilled POs.json",
            unfulfilled_json_path="output/Unfulfilled POs.json",
        )
        pipeline.run(tracker_path="MAXISMED COMMISSION TRACKER.xlsm")
    """

    SEQNO_KEY_INDEX = 6  # position of SeqNo within the pipe-delimited JSON key
    SCRUBSHEET_DISCREPANCY_TEXT = "SCRUBSHEET DISCREPANCY"
    DUPLICATE_COMMENT_TEXT = "Duplicate PO#"
    COMMENT_AUTHOR = "PO Pipeline"
    DUPLICATE_ROW_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    UNFULFILLED_ROW_FILL = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

    def __init__(
        self,
        fulfilled_json_path: Path | str,
        unfulfilled_json_path: Path | str,
        email_user: str | None = None,
        gmail_app_password: str | None = None,
        sender_filter: str | None = None,
        subject_filter: str | None = None,
        po_column_name: str | None = None,
        key_columns: list[str] | None = None,
        sheet_name: str = "Invoice Line Items",
        seqno_col: int = 7,
        po_col: int = 8,
        otp_sender_filter: str = "messaging.microsoft.com",
        otp_subject_hints: tuple[str, ...] = ("code", "one-time", "one time", "password", "temporary"),
        otp_wait_timeout_seconds: int = 180,
        otp_poll_interval_seconds: int = 3,
        download_dir: Path | str = "downloads",
        headless: bool = False,
        imap_host: str = "imap.gmail.com",
    ):
        self.fulfilled_json_path = Path(fulfilled_json_path)
        self.unfulfilled_json_path = Path(unfulfilled_json_path)
        self.email_user = email_user
        self.gmail_app_password = gmail_app_password
        self.sender_filter = sender_filter
        self.subject_filter = subject_filter
        self.po_column_name = po_column_name
        self.key_columns = key_columns
        self.sheet_name = sheet_name
        self.seqno_col = seqno_col
        self.po_col = po_col
        self.otp_sender_filter = otp_sender_filter
        self.otp_subject_hints = otp_subject_hints
        self.otp_wait_timeout_seconds = otp_wait_timeout_seconds
        self.otp_poll_interval_seconds = otp_poll_interval_seconds
        self.download_dir = Path(download_dir)
        self.headless = headless
        self.imap_host = imap_host

        self._imap: imaplib.IMAP4_SSL | None = None

    # -----------------------------------------------------------------
    # Public API
    # -----------------------------------------------------------------
    def run(self, tracker_path: Path | str) -> None:
        """Fetches new PO emails, merges/splits the JSON records, and backfills the tracker."""
        new_records = self.fetch_emails()
        fulfilled, unfulfilled = self.merge_and_split(new_records)
        self.backfill_tracker(fulfilled, unfulfilled, tracker_path)

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, tb):
        self._close()

    # ===================================================================
    # Stage 1: email fetching / OME redemption / spreadsheet parsing
    # ===================================================================
    def fetch_emails(self) -> dict[str, str]:
        """Process every unread matching email and return the combined dataset."""
        print("[gmail] connecting via IMAP...", flush=True)
        self._connect()
        print("[gmail] connected", flush=True)

        all_records: dict[str, str] = {}
        try:
            with sync_playwright() as p:
                browser = p.chromium.launch(headless=self.headless)
                try:
                    for uid, msg in self._iter_matching_messages():
                        print(f"[email] processing uid={uid.decode()}: {msg.get('Subject', '')!r}", flush=True)
                        read_link = self._extract_read_message_link(self._get_html_body(msg))
                        page = self._redeem_encrypted_message(browser, read_link)
                        try:
                            xlsx_path = self._download_attachment(page)
                        except Exception:
                            self._save_debug_snapshot(page)
                            raise
                        all_records.update(self._parse_spreadsheet(xlsx_path))
                        page.close()
                        self._imap.uid("store", uid, "+FLAGS", "\\Seen")
                finally:
                    browser.close()
        finally:
            self._close()

        return all_records

    # -----------------------------------------------------------------
    # IMAP connection lifecycle
    # -----------------------------------------------------------------
    def _connect(self):
        # Without an explicit socket timeout, a stalled connection blocks
        # forever instead of raising -- this turns that into a catchable error.
        self._imap = imaplib.IMAP4_SSL(self.imap_host, timeout=30)
        self._imap.login(self.email_user, self.gmail_app_password)
        self._imap.select("INBOX")

    def _close(self):
        if self._imap is not None:
            try:
                self._imap.logout()
            except Exception:
                pass
            self._imap = None

    # -----------------------------------------------------------------
    # Gmail / IMAP side
    # -----------------------------------------------------------------
    @staticmethod
    def _has_attachment(msg: email.message.Message) -> bool:
        for part in msg.walk():
            if "attachment" in str(part.get("Content-Disposition", "")).lower():
                return True
        return False

    @staticmethod
    def _get_html_body(msg: email.message.Message) -> str | None:
        if msg.is_multipart():
            for part in msg.walk():
                if part.get_content_type() == "text/html":
                    charset = part.get_content_charset() or "utf-8"
                    return part.get_payload(decode=True).decode(charset, errors="replace")
        elif msg.get_content_type() == "text/html":
            charset = msg.get_content_charset() or "utf-8"
            return msg.get_payload(decode=True).decode(charset, errors="replace")
        return None

    @staticmethod
    def _extract_read_message_link(html_body: str) -> str:
        soup = BeautifulSoup(html_body, "html.parser")
        candidates = []
        for a in soup.find_all("a", href=True):
            href = a["href"]
            text = a.get_text(separator=" ", strip=True).replace("\xa0", " ")
            host = urlparse(href).netloc.lower()
            href_ok = any(
                host == suffix or host.endswith("." + suffix) for suffix in READ_LINK_HREF_HOST_SUFFIXES
            )
            if READ_LINK_PATTERN.search(text) and href_ok:
                candidates.append(href)
        if not candidates:
            raise RuntimeError("Could not find a 'Read the message' link in the email body")
        if len(candidates) > 1:
            raise RuntimeError(
                f"Found {len(candidates)} candidate 'Read the message' links "
                "(expected exactly one) -- inspect the email body, it may contain "
                "a quoted older notification"
            )
        return candidates[0]

    @staticmethod
    def _extract_otp_code(body_html: str) -> str | None:
        soup = BeautifulSoup(body_html, "html.parser")
        span = soup.find(id=re.compile(r"(passcode|password|key|code)$", re.I))
        if span:
            text = span.get_text(strip=True)
            if text.isdigit():
                return text
        match = re.search(r"\b\d{6,8}\b", soup.get_text())
        return match.group(0) if match else None

    def _fetch_message_by_uid(self, uid: bytes) -> email.message.Message | None:
        """
        Fetches a message body by IMAP UID rather than sequence number.
        Sequence numbers are just a message's position in the mailbox and
        shift whenever the mailbox changes (new mail arriving, messages
        moved/expunged) -- even mid-run. That instability caused a real
        crash here: a stale sequence number's FETCH came back without a
        body literal at all, and unconditionally indexing into the response
        (msg_data[0][1]) blew up trying to .decode() something that wasn't
        bytes. UIDs stay fixed for the life of the message, so search/fetch/
        store all use self._imap.uid(...) instead of the bare equivalents.
        A malformed/empty response is now just skipped instead of crashing.
        """
        typ, msg_data = self._imap.uid("fetch", uid, "(BODY.PEEK[])")
        if typ != "OK" or not msg_data or not isinstance(msg_data[0], tuple):
            return None
        return email.message_from_bytes(msg_data[0][1])

    def _iter_matching_messages(self):
        """Yields (uid, email.message.Message) for each unread, matching, PO-request email."""
        typ, data = self._imap.uid(
            "search",
            None,
            "UNSEEN",
            "FROM",
            f'"{self.sender_filter}"',
            "HEADER",
            "SUBJECT",
            f'"{self.subject_filter}"',
        )
        if typ != "OK":
            raise RuntimeError(f"IMAP search failed: {data}")

        uids = data[0].split()
        print(f"[gmail] {len(uids)} unread candidate email(s) found", flush=True)

        for uid in uids:
            # BODY.PEEK[] fetches without marking \Seen -- we only mark it read
            # after successfully downloading & parsing its attachment, so a
            # mid-run failure leaves the email unread and retryable next run.
            msg = self._fetch_message_by_uid(uid)
            if msg is None:
                continue
            if self._has_attachment(msg):
                yield uid, msg

    def _wait_for_otp_code(self, exclude_uids: set[bytes]) -> str:
        print("[otp] waiting for one-time-passcode email...", flush=True)
        deadline = time.time() + self.otp_wait_timeout_seconds
        while time.time() < deadline:
            self._imap.noop()  # refresh IMAP's view of the mailbox
            typ, data = self._imap.uid("search", None, "UNSEEN", "FROM", f'"{self.otp_sender_filter}"')
            if typ == "OK":
                # Skip UIDs that already existed before we requested this passcode --
                # otherwise a stale, expired OTP email from an earlier attempt can be
                # picked up instead of the fresh one we just triggered.
                new_uids = [uid for uid in reversed(data[0].split()) if uid not in exclude_uids]
                for uid in new_uids:  # newest first
                    msg = self._fetch_message_by_uid(uid)
                    if msg is None:
                        continue
                    subject = msg.get("Subject", "").lower()
                    if not any(hint in subject for hint in self.otp_subject_hints):
                        continue
                    body = self._get_html_body(msg) or ""
                    if not body:
                        payload = msg.get_payload(decode=True)
                        if payload:
                            body = payload.decode(msg.get_content_charset() or "utf-8", errors="replace")
                    code = self._extract_otp_code(body)
                    if code:
                        self._imap.uid("store", uid, "+FLAGS", "\\Seen")
                        print(f"[otp] code received: {code}", flush=True)
                        return code
            time.sleep(self.otp_poll_interval_seconds)
        raise TimeoutError("Timed out waiting for the one-time-passcode email to arrive")

    # -----------------------------------------------------------------
    # Playwright / Microsoft OME side
    # -----------------------------------------------------------------
    @staticmethod
    def _invoke_click(locator):
        # Calls the element's native .click() directly via JS instead of
        # Playwright's default click, which moves a simulated pointer to
        # computed coordinates and dispatches real mouse events. No mouse
        # movement, no coordinates.
        locator.evaluate("el => el.click()")

    @staticmethod
    def _locate_in_page_or_frames(page, locator_fn, timeout_ms: int = 45000):
        # .count() checks the DOM instantly with no retry, which fails during
        # transitional states (e.g. the "Logging you in..." interstitial
        # between OTP verification and the real message loading). Poll instead.
        deadline = time.time() + timeout_ms / 1000
        while True:
            loc = locator_fn(page)
            if loc.count() > 0:
                return loc.first
            for frame in page.frames:
                loc = locator_fn(frame)
                if loc.count() > 0:
                    return loc.first
            if time.time() >= deadline:
                raise RuntimeError("Could not locate element on the page or in any of its frames (timed out)")
            page.wait_for_timeout(500)

    def _redeem_encrypted_message(self, browser, read_link: str):
        page = browser.new_page()
        print(f"[ome] opening {read_link[:80]}...", flush=True)
        page.goto(read_link)

        # Must be a UID search, matching _wait_for_otp_code's UID search below --
        # comparing UIDs against sequence numbers (or vice versa) would make the
        # exclude_uids filter meaningless.
        typ, data = self._imap.uid("search", None, "FROM", f'"{self.otp_sender_filter}"')
        existing_uids = set(data[0].split()) if typ == "OK" else set()

        self._invoke_click(page.get_by_role("button", name=OTP_LABEL_RE))
        print("[ome] requested a one-time passcode", flush=True)

        otp_code = self._wait_for_otp_code(exclude_uids=existing_uids)

        page.get_by_label(OTP_LABEL_RE).fill(otp_code)
        self._invoke_click(page.get_by_text("Continue", exact=True))

        # networkidle is a weak signal here -- it's trivially satisfied if the
        # click did nothing at all. Confirm the OTP screen actually went away.
        page.get_by_text(OTP_SENT_NOTICE_RE).wait_for(state="detached", timeout=30000)
        print("[ome] message unlocked", flush=True)
        return page

    def _download_attachment(self, page) -> Path:
        print("[ome] looking for the attachment...", flush=True)

        # The page has multiple ChevronDown icons (e.g. Reply All's dropdown up
        # in the header), so grabbing the first one on the page is ambiguous.
        # Anchor on the attachment filename text instead and take the nearest
        # chevron that appears AFTER it in document order -- the Reply All
        # chevron comes before the attachment row, so this reliably skips it.
        filename_text = self._locate_in_page_or_frames(
            page, lambda ctx: ctx.get_by_text(re.compile(r"\.xlsx$", re.I))
        )
        chevron = filename_text.locator("xpath=following::*[@data-icon-name='ChevronDown'][1]")
        self._invoke_click(chevron)
        print("[ome] opened attachment menu", flush=True)

        download_item = self._locate_in_page_or_frames(
            page, lambda ctx: ctx.get_by_text("Download", exact=True)
        )
        with page.expect_download() as dl_info:
            self._invoke_click(download_item)
        download = dl_info.value

        self.download_dir.mkdir(exist_ok=True)
        dest = self.download_dir / download.suggested_filename
        download.save_as(str(dest))
        print(f"[ome] downloaded {dest}", flush=True)
        return dest

    @staticmethod
    def _save_debug_snapshot(page):
        try:
            debug_dir = Path("debug")
            debug_dir.mkdir(exist_ok=True)
            page.screenshot(path=str(debug_dir / "attachment_page.png"), full_page=True)
            (debug_dir / "attachment_page.html").write_text(page.content(), encoding="utf-8")
        except Exception as capture_error:
            print(f"[debug] could not capture debug info: {capture_error}", flush=True)

    def _parse_spreadsheet(self, path: Path) -> dict[str, str]:
        df = pd.read_excel(path)
        df = df[df[self.po_column_name].notna() & (df[self.po_column_name].astype(str).str.strip() != "")]

        records: dict[str, str] = {}
        for _, row in df.iterrows():
            key = "|".join(str(row[col]) for col in self.key_columns)
            records[key] = str(row[self.po_column_name])
        return records

    # ===================================================================
    # Stage 2: JSON merge + split
    # ===================================================================
    @staticmethod
    def _load_json(path: Path) -> dict[str, str]:
        return json.loads(path.read_text()) if path.exists() else {}

    @staticmethod
    def _is_numeric(value: str) -> bool:
        try:
            float(value.strip())
        except ValueError:
            return False
        return True

    def merge_and_split(self, new_records: dict[str, str]) -> tuple[dict[str, str], dict[str, str]]:
        """
        Combines new_records with whatever's already on disk across both the
        fulfilled and unfulfilled files (new wins on duplicate keys -- a SeqNo
        that was previously unfulfilled and now has a real PO moves over to
        fulfilled), re-splits on whether the PO value is numeric, and writes
        both files back out.
        """
        merged = self._load_json(self.fulfilled_json_path)
        merged.update(self._load_json(self.unfulfilled_json_path))
        merged.update(new_records)

        fulfilled = {k: v for k, v in merged.items() if self._is_numeric(v)}
        unfulfilled = {k: v for k, v in merged.items() if not self._is_numeric(v)}

        self.fulfilled_json_path.parent.mkdir(exist_ok=True, parents=True)
        self.fulfilled_json_path.write_text(json.dumps(fulfilled, indent=2))
        self.unfulfilled_json_path.write_text(json.dumps(unfulfilled, indent=2))

        print(f"Fulfilled POs: {len(fulfilled)} records -> {self.fulfilled_json_path}")
        print(f"Unfulfilled POs: {len(unfulfilled)} records -> {self.unfulfilled_json_path}")

        self._write_csv(fulfilled, self.fulfilled_json_path.with_suffix(".csv"))
        self._write_csv(unfulfilled, self.unfulfilled_json_path.with_suffix(".csv"))
        return fulfilled, unfulfilled

    def _write_csv(self, records: dict[str, str], path: Path) -> None:
        """Writes a two-column SeqNo,PO CSV (for the VBA backfill) from a records dict."""
        rows = 0
        with open(path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["SeqNo", "PO"])
            for key, po in records.items():
                parts = key.split("|")
                if len(parts) <= self.SEQNO_KEY_INDEX:
                    continue
                seqno = parts[self.SEQNO_KEY_INDEX]
                if not seqno or seqno == "nan":
                    continue
                writer.writerow([seqno, po])
                rows += 1
        print(f"CSV: {rows} rows -> {path}")

    # ===================================================================
    # Stage 3: tracker backfill
    # ===================================================================
    def _build_seqno_map(self, records: dict[str, str]) -> dict[str, str]:
        seqno_map: dict[str, str] = {}
        for key, value in records.items():
            parts = key.split("|")
            if len(parts) <= self.SEQNO_KEY_INDEX:
                continue
            seqno = parts[self.SEQNO_KEY_INDEX]
            if not seqno or seqno == "nan":
                continue
            if seqno in seqno_map and seqno_map[seqno] != value:
                raise ValueError(
                    f"Duplicate SeqNo error: {seqno!r} appears twice with conflicting "
                    f"values {seqno_map[seqno]!r} and {value!r}"
                )
            seqno_map[seqno] = value
        return seqno_map

    @staticmethod
    def _as_po_value(value: str):
        stripped = value.strip()
        return int(stripped) if stripped.isdigit() else stripped

    @staticmethod
    def _normalize_po(value) -> str | None:
        """Canonical form for PO comparison -- 181114, "181114", and 181114.0
        (int, text, and float-formatted variants that show up in the sheet)
        all need to compare equal."""
        if value in (None, ""):
            return None
        text = str(value).strip()
        try:
            return str(int(float(text)))
        except ValueError:
            return text

    def _fill_row(self, ws, row: int, fill: PatternFill) -> None:
        for col in range(1, ws.max_column + 1):
            ws.cell(row=row, column=col).fill = fill

    def backfill_tracker(
        self, fulfilled: dict[str, str], unfulfilled: dict[str, str], tracker_path: Path | str
    ) -> None:
        """
        Backfills missing PO numbers (column H) on the tracker sheet at
        tracker_path, for rows whose SeqNo (column G) starts with "IF" and
        whose PO cell is currently blank.

        Behavior per eligible row:
          - SeqNo found in `fulfilled`   -> write the PO number into H. If
            that PO number is already present on another "IF-" row (either
            pre-existing before this run, or inserted earlier in this same
            run), THIS row also gets a "Duplicate PO#" comment and its
            whole row highlighted. Only rows filled during this run are
            ever flagged -- a row that already had a PO before this run
            started is never rescanned or flagged, even if it happens to
            share a PO number with something else.
          - SeqNo found in `unfulfilled` -> write the literal text
            "SCRUBSHEET DISCREPANCY" into H, highlight the whole row red,
            and attach a cell comment on H with the original note text
            (e.g. "Mt total is $7,746") so the underlying detail is still
            one hover away
          - SeqNo not found in either    -> leave untouched (PO simply
            hasn't been sent yet -- not an error)

        Never overwrites a PO cell that already has a value.
        """
        tracker_path = Path(tracker_path)
        fulfilled_map = self._build_seqno_map(fulfilled)
        unfulfilled_map = self._build_seqno_map(unfulfilled)

        wb = openpyxl.load_workbook(tracker_path, data_only=False, keep_vba=True)
        ws = wb[self.sheet_name]

        # Baseline of PO values already on IF- rows before this run's
        # insertions -- scoped to IF- rows only, since other vendors
        # legitimately reuse one PO number across a batch of same-day line
        # items, which isn't an error. Used to catch a freshly-inserted PO
        # colliding with something that was already there.
        existing_pos: dict[str, int] = {}
        for row in range(2, ws.max_row + 1):
            seqno = ws.cell(row=row, column=self.seqno_col).value
            if not seqno or not str(seqno).startswith("IF"):
                continue
            norm = self._normalize_po(ws.cell(row=row, column=self.po_col).value)
            if norm is not None:
                existing_pos.setdefault(norm, row)

        filled = []
        flagged = []
        duplicates = []
        for row in range(2, ws.max_row + 1):
            seqno_cell = ws.cell(row=row, column=self.seqno_col)
            po_cell = ws.cell(row=row, column=self.po_col)

            seqno = seqno_cell.value
            if not seqno or not str(seqno).startswith("IF"):
                continue
            if po_cell.value not in (None, ""):
                continue  # never overwrite an existing PO

            seqno = str(seqno)
            if seqno in fulfilled_map:
                po_value = self._as_po_value(fulfilled_map[seqno])
                po_cell.value = po_value
                filled.append((row, seqno, po_value))

                norm = self._normalize_po(po_value)
                existing_row = existing_pos.get(norm)
                if existing_row is not None and existing_row != row:
                    po_cell.comment = Comment(self.DUPLICATE_COMMENT_TEXT, self.COMMENT_AUTHOR)
                    self._fill_row(ws, row, self.DUPLICATE_ROW_FILL)
                    duplicates.append((row, seqno, po_value, existing_row))
                existing_pos.setdefault(norm, row)
            elif seqno in unfulfilled_map:
                po_cell.value = self.SCRUBSHEET_DISCREPANCY_TEXT
                po_cell.comment = Comment(unfulfilled_map[seqno], self.COMMENT_AUTHOR)
                self._fill_row(ws, row, self.UNFULFILLED_ROW_FILL)
                flagged.append((row, seqno))
            # else: no JSON record yet -- leave untouched, not an error

        wb.save(tracker_path)

        print(f"Filled {len(filled)} PO(s):")
        for row, seqno, po in filled:
            print(f"  - row {row} ({seqno}): PO = {po}")
        print(f"Flagged {len(flagged)} row(s) as '{self.SCRUBSHEET_DISCREPANCY_TEXT}':")
        for row, seqno in flagged:
            print(f"  - row {row} ({seqno})")
        print(f"Flagged {len(duplicates)} row(s) as '{self.DUPLICATE_COMMENT_TEXT}':")
        for row, seqno, po, existing_row in duplicates:
            print(f"  - row {row} ({seqno}): PO {po} already used on row {existing_row}")
